"""DOCX hyperlinks and inline formatting: tags on parse, rebuilt runs and links on reconstruct."""

from __future__ import annotations

import io

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from parsers import docx_parser

URL = "https://example.com/terms"


def _add_link(para, text: str, bold: bool = False):
    r_id = para.part.relate_to(URL, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rpr.append(style)
    if bold:
        rpr.append(OxmlElement("w:b"))
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    para._p.append(link)


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _bold_and_link_docx(link_bold: bool = False) -> bytes:
    """One paragraph: "Read the " + bold "terms" + " at " + link "our site" + " today."."""
    doc = Document()
    para = doc.add_paragraph("Read the ")
    para.add_run("terms").bold = True
    para.add_run(" at ")
    _add_link(para, "our site", bold=link_bold)
    para.add_run(" today.")
    return _save(doc)


def _link_only_docx() -> bytes:
    """One paragraph with a link and no other formatting."""
    doc = Document()
    para = doc.add_paragraph("Visit ")
    _add_link(para, "our site")
    para.add_run(" today.")
    return _save(doc)


def _first_paragraph(content: bytes):
    return Document(io.BytesIO(content)).paragraphs[0]


def _runs_in_order(para):
    """(text, bold, inside a hyperlink) for every run, links included."""
    out = []
    for item in para.iter_inner_content():
        if hasattr(item, "address"):
            out.extend((r.text, bool(r.bold), True) for r in item.runs)
        else:
            out.append((item.text, bool(item.bold), False))
    return out


# --- parse ---------------------------------------------------------------


def test_parse_formatted_paragraph_keeps_link_text_as_tag():
    segments = docx_parser.parse(_bold_and_link_docx())

    assert len(segments) == 1
    assert segments[0].text == "Read the terms at our site today."
    assert segments[0].rich_text == (
        'Read the <g id="1" ctype="bold">terms</g> at <g id="2" ctype="link">our site</g> today.'
    )


def test_parse_link_without_other_formatting_is_tagged():
    segments = docx_parser.parse(_link_only_docx())

    assert segments[0].text == "Visit our site today."
    assert segments[0].rich_text == 'Visit <g id="1" ctype="link">our site</g> today.'


def test_parse_formatting_inside_link_nests_in_link_tag():
    segments = docx_parser.parse(_bold_and_link_docx(link_bold=True))

    assert segments[0].rich_text == (
        'Read the <g id="1" ctype="bold">terms</g> at '
        '<g id="2" ctype="link"><g id="3" ctype="bold">our site</g></g> today.'
    )


# --- reconstruct ---------------------------------------------------------


def test_reconstruct_rebuilds_link_and_bold_from_tags():
    out = docx_parser.reconstruct(
        _bold_and_link_docx(),
        {"para:0": 'Lea los <g id="1" ctype="bold">términos</g> en <g id="2" ctype="link">nuestro sitio</g> hoy.'},
    )

    para = _first_paragraph(out)
    assert para.text == "Lea los términos en nuestro sitio hoy."
    assert [(h.text, h.address) for h in para.hyperlinks] == [("nuestro sitio", URL)]
    assert _runs_in_order(para) == [
        ("Lea los ", False, False),
        ("términos", True, False),
        (" en ", False, False),
        ("nuestro sitio", False, True),
        (" hoy.", False, False),
    ]


def test_reconstruct_link_keeps_its_run_style():
    out = docx_parser.reconstruct(
        _link_only_docx(), {"para:0": 'Visite <g id="1" ctype="link">nuestro sitio</g> hoy.'}
    )

    link_run = _first_paragraph(out).hyperlinks[0].runs[0]
    # The blank template defines no "Hyperlink" style, so read the reference itself.
    assert link_run._r.rPr.rStyle.val == "Hyperlink"


def test_reconstruct_bold_inside_link():
    out = docx_parser.reconstruct(
        _bold_and_link_docx(link_bold=True),
        {"para:0": 'Lea los <g id="1" ctype="bold">términos</g> en <g id="2" ctype="link"><g id="3" ctype="bold">nuestro sitio</g></g> hoy.'},
    )

    para = _first_paragraph(out)
    assert para.text == "Lea los términos en nuestro sitio hoy."
    assert ("nuestro sitio", True, True) in _runs_in_order(para)


def test_reconstruct_moved_tags_follow_the_translation():
    """The translation decides where each tag sits, not the source order."""
    out = docx_parser.reconstruct(
        _bold_and_link_docx(),
        {"para:0": '<g id="2" ctype="link">Unser Portal</g> listet die <g id="1" ctype="bold">Bedingungen</g>.'},
    )

    para = _first_paragraph(out)
    assert para.text == "Unser Portal listet die Bedingungen."
    assert [(h.text, h.address) for h in para.hyperlinks] == [("Unser Portal", URL)]
    assert ("Bedingungen", True, False) in _runs_in_order(para)


def test_reconstruct_plain_translation_replaces_link_text():
    """A translation without tags replaces the whole paragraph, link text included."""
    out = docx_parser.reconstruct(_link_only_docx(), {"para:0": "Visite nuestro sitio hoy."})

    para = _first_paragraph(out)
    assert para.text == "Visite nuestro sitio hoy."
    assert para.hyperlinks == []


def test_reconstruct_dropped_link_tag_gives_plain_text():
    out = docx_parser.reconstruct(
        _bold_and_link_docx(), {"para:0": 'Lea los <g id="1" ctype="bold">términos</g> hoy.'}
    )

    para = _first_paragraph(out)
    assert para.text == "Lea los términos hoy."
    assert para.hyperlinks == []


def test_reconstruct_ignores_stray_and_unknown_tags():
    out = docx_parser.reconstruct(
        _link_only_docx(),
        {"para:0": 'Visite </g><g id="9" ctype="link">nuestro sitio</g> <x id="7" ctype="ph"/>hoy.'},
    )

    para = _first_paragraph(out)
    assert para.text == "Visite nuestro sitio hoy."
    # Tag 9 is not in the source, so there is no link to rebuild.
    assert para.hyperlinks == []


def test_reconstruct_line_break_tag_becomes_break():
    out = docx_parser.reconstruct(
        _link_only_docx(), {"para:0": 'Visite<x id="2" ctype="lb" original_tag="br"/><g id="1" ctype="link">nuestro sitio</g>'}
    )

    para = _first_paragraph(out)
    assert para.text == "Visite\nnuestro sitio"
    assert para._p.xpath(".//w:br")


def test_reconstruct_untagged_bold_paragraph_stays_bold():
    """A paragraph in one bold run has no tags; its translation keeps the run's formatting."""
    doc = Document()
    doc.add_paragraph().add_run("Warning").bold = True
    out = docx_parser.reconstruct(_save(doc), {"para:0": "Advertencia"})

    assert _runs_in_order(_first_paragraph(out)) == [("Advertencia", True, False)]


def test_reconstruct_text_outside_bold_tag_is_not_bold():
    """When every source run is tagged, text the translator puts outside the tags is plain."""
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Bold").bold = True
    para.add_run("Italic").italic = True
    out = docx_parser.reconstruct(
        _save(doc), {"para:0": 'Plain <g id="1" ctype="bold">Fett</g><g id="2" ctype="italic">Kursiv</g>'}
    )

    runs = _first_paragraph(out).runs
    assert [(r.text, bool(r.bold), bool(r.italic)) for r in runs] == [
        ("Plain ", False, False),
        ("Fett", True, False),
        ("Kursiv", False, True),
    ]


def test_parse_reconstruct_round_trip_keeps_rich_text():
    """Writing the source rich text back gives a file that parses to the same rich text."""
    original = _bold_and_link_docx(link_bold=True)
    before = docx_parser.parse(original)[0]

    after = docx_parser.parse(docx_parser.reconstruct(original, {"para:0": before.rich_text}))[0]

    assert after.text == before.text
    assert after.rich_text == before.rich_text
