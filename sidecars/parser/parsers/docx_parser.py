"""DOCX parser using python-docx. Extracts paragraphs from body, tables, headers, and footers."""

from __future__ import annotations

import io
from collections.abc import Iterator
from dataclasses import dataclass

from docx import Document
from docx.text.paragraph import Paragraph

from parsers.docx_inline import paragraph_tags, write_translation


@dataclass
class Segment:
    key: str
    text: str
    position: int
    context: str  # "body", "table", "header", "footer"
    rich_text: str = ""  # XLIFF <g> tags for bold, italic and links; see docx_inline


def _paragraphs(doc) -> Iterator[tuple[Paragraph, str]]:
    """Every paragraph with text, with its context, in segment-key order.

    parse and reconstruct both walk this, so a key names the same paragraph in both.
    """
    def with_text(paras, ctx):
        for para in paras:
            if para.text.strip():
                yield para, ctx

    yield from with_text(doc.paragraphs, "body")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from with_text(cell.paragraphs, "table")
    for section in doc.sections:
        yield from with_text(section.header.paragraphs, "header")
        yield from with_text(section.footer.paragraphs, "footer")


def parse(content: bytes) -> list[Segment]:
    """Extract translatable segments from a DOCX file."""
    doc = Document(io.BytesIO(content))
    return [
        Segment(
            key=f"para:{pos}",
            text=para.text.strip(),
            position=pos,
            context=ctx,
            rich_text=paragraph_tags(para).rich_text,
        )
        for pos, (para, ctx) in enumerate(_paragraphs(doc))
    ]


def reconstruct(original: bytes, translations: dict[str, str]) -> bytes:
    """Rebuild a DOCX file with translations replacing original text.

    A translation may carry the rich text tags parse emits; write_translation
    turns them back into formatted runs and links.
    """
    doc = Document(io.BytesIO(original))
    # Collect first: writing a translation must not shift the walk.
    targets = [(f"para:{pos}", para) for pos, (para, _) in enumerate(_paragraphs(doc))]
    for key, para in targets:
        if key in translations:
            write_translation(para, translations[key])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
